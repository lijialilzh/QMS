import io
import logging
from sqlalchemy import select, delete, func, or_
from sqlalchemy.sql import asc
from openpyxl import load_workbook
from ..model.prod_hospital import ProdHospital
from ..obj.tobj_prod_hospital import ProdHospitalForm
from ..obj.vobj_prod_hospital import ProdHospitalObj
from ..utils.sql_ctx import db
from ..utils.i18n import ts
from ..obj import Page, Resp
from . import msg_err_db

logger = logging.getLogger(__name__)

_HEADER_MAP = {
    "医院名称": "org_name",
    "对方单位名称": "org_name",
    "区域划分": "region",
    "区域": "region",
    "医院编号": "hospital_no",
    "省份": "province",
    "城市信息": "city",
    "城市": "city",
    "合同编号": "contract_no",
}


def _cell_str(value):
    if value is None:
        return ""
    if isinstance(value, float):
        if value == int(value):
            return str(int(value))
        return str(value).strip()
    return str(value).strip()


def _read_excel_grid(file_bytes: bytes):
    if not file_bytes:
        return []
    if file_bytes[:2] == b"PK":
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
        ws = wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    import xlrd
    book = xlrd.open_workbook(file_contents=file_bytes)
    sheet = book.sheet_by_index(0)
    return [sheet.row_values(i) for i in range(sheet.nrows)]


def _header_index(grid):
    for i, row in enumerate(grid or []):
        cells = [_cell_str(c) for c in (row or [])]
        mapped = [_HEADER_MAP.get(c) for c in cells]
        if "org_name" in mapped and "hospital_no" in mapped:
            return i, cells
    return None, None


class Server(object):

    async def add_prod_hospital(self, form: ProdHospitalForm):
        try:
            if form.prod_id is None:
                form.prod_id = 0
            row = ProdHospital(**form.dict())
            row.id = None
            row.prod_id = form.prod_id or 0
            db.session.add(row)
            db.session.commit()
            return Resp.resp_ok()
        except Exception:
            logger.exception("")
            db.session.rollback()
        return Resp.resp_err(msg=ts(msg_err_db))

    async def update_prod_hospital(self, form: ProdHospitalForm):
        try:
            sql = select(ProdHospital).where(ProdHospital.id == form.id)
            row: ProdHospital = db.session.execute(sql).scalars().first()
            if not row:
                return Resp.resp_err(msg=ts("msg_obj_null"))
            for key, value in form.dict().items():
                if key in ("id", "prod_id"):
                    continue
                setattr(row, key, value)
            db.session.commit()
            return Resp.resp_ok()
        except Exception:
            logger.exception("")
            db.session.rollback()
        return Resp.resp_err(msg=ts(msg_err_db))

    async def delete_prod_hospitals(self, ids: list):
        ids = [int(i) for i in (ids or []) if str(i).strip().isdigit()]
        if not ids:
            return Resp.resp_err(msg=ts("msg_err_param"))
        db.session.execute(delete(ProdHospital).where(ProdHospital.id.in_(ids)))
        db.session.commit()
        return Resp.resp_ok()

    async def list_prod_hospital(self, prod_id: int = None, fuzzy: str = None,
                                 page_index: int = 0, page_size: int = 10):
        page_index = page_index if page_index >= 0 else 0
        page_size = page_size if page_size > 0 else 10
        # 全局清单：不按产品过滤
        sql = select(ProdHospital)
        kw = (fuzzy or "").strip()
        if kw:
            like = f"%{kw}%"
            sql = sql.where(or_(
                ProdHospital.contract_no.like(like),
                ProdHospital.org_name.like(like),
                ProdHospital.hospital_no.like(like),
                ProdHospital.region.like(like),
                ProdHospital.province.like(like),
                ProdHospital.city.like(like),
            ))

        sql_count = select(func.count()).select_from(sql)
        total = db.session.execute(sql_count).scalars().first()

        sql = sql.order_by(asc(ProdHospital.sort_order), asc(ProdHospital.id))
        sql = sql.offset(page_size * page_index).limit(page_size)
        rows: list[ProdHospital] = db.session.execute(sql).scalars().all()
        objs = [ProdHospitalObj(**row.dict()) for row in rows]
        return Resp.resp_ok(data=Page(total=total, page_size=page_size, rows=objs, page_index=page_index))

    async def import_prod_hospitals(self, prod_id: int, file_bytes: bytes, replace: bool = True):
        try:
            grid = _read_excel_grid(file_bytes)
            header_idx, header_cells = _header_index(grid)
            if header_idx is None:
                return Resp.resp_err(msg=ts("msg_err_param"))
            col_map = {}
            for i, name in enumerate(header_cells):
                field = _HEADER_MAP.get(name)
                if field and field not in col_map:
                    col_map[field] = i
            if replace:
                db.session.execute(delete(ProdHospital))
            imported = 0
            sort_order = 0
            seen_nos = set()
            for row in grid[header_idx + 1:]:
                def cell(field):
                    i = col_map.get(field)
                    if i is None or i >= len(row):
                        return ""
                    return _cell_str(row[i])
                org_name = cell("org_name")
                hospital_no = cell("hospital_no")
                if not (org_name or hospital_no):
                    continue
                no_key = hospital_no.upper()
                if no_key and no_key in seen_nos:
                    continue
                if no_key:
                    seen_nos.add(no_key)
                sort_order += 1
                db.session.add(ProdHospital(
                    prod_id=0,
                    contract_no=cell("contract_no"),
                    org_name=org_name,
                    hospital_no=hospital_no,
                    region=cell("region"),
                    province=cell("province"),
                    city=cell("city"),
                    sort_order=sort_order,
                ))
                imported += 1
            db.session.commit()
            return Resp.resp_ok(data={"imported": imported})
        except Exception:
            logger.exception("")
            db.session.rollback()
        return Resp.resp_err(msg=ts(msg_err_db))
