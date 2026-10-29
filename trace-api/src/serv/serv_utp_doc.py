#!/usr/bin/env python
# encoding: utf-8

# 用户测试计划服务层（开发文件，PDP 风格章节树）。默认内容取自 src-res/utp_default_content.json。

import base64
import copy
import json
import logging
import os
import re
from io import BytesIO
from typing import List

from sqlalchemy import delete, func, select
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT

from ..model.product import Product
from ..model.utp_doc import UtpDoc
from ..model.prod_dhf import ProdDhf
from ..model.project_member import ProjectMember
from .serv_prod_runtime_env import apply_runtime_to_pdp_content
from ..obj import Page, Resp
from ..obj.tobj_role import Roles
from ..obj.vobj_user import UserObj
from ..obj.tobj_utp_doc import UtpDocForm
from ..obj.vobj_utp_doc import UtpDocObj
from ..utils.i18n import ts
from ..utils.sql_ctx import db
from . import msg_err_db
from . import serv_review_util
from .serv_utils import new_version, sync_file_no_version
from .serv_utils import docx_util
from .serv_product import product_scope_filter

logger = logging.getLogger(__name__)

_DEFAULT_FILE = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "src-res", "utp_default_content.json")
try:
    with open(_DEFAULT_FILE, encoding="utf-8") as _f:
        DEFAULT_UTP_CONTENT = json.load(_f)
except Exception:
    DEFAULT_UTP_CONTENT = {"sections": []}

# 模板基准产品名（用于全文替换为当前产品名）。
BASE_NAME = "InferOperate Suite"
DOC_NAME = "用户测试计划"
DOC_KEY = "utp"


class Server(object):

    def __normalize_node(self, node):
        if not isinstance(node, dict):
            return {"title": str(node or ""), "body": "", "tables": [], "children": []}
        result = dict(node)
        result["title"] = str(result.get("title") or "")
        result["body"] = str(result.get("body") or "")
        tables = result.get("tables")
        if not isinstance(tables, list):
            tables = []
        norm_tables = []
        for table in tables:
            if isinstance(table, list):
                norm_tables.append([[str(c) if c is not None else "" for c in (row or [])] for row in table if isinstance(row, list)])
        result["tables"] = norm_tables
        # 兼容旧数据：如果有 blocks（之前改造残留），把 blocks 中的 table/text 提取回 tables/body
        blocks = result.get("blocks")
        if isinstance(blocks, list) and blocks and not result["body"] and not norm_tables:
            parts = []
            tbls = []
            for blk in blocks:
                if not isinstance(blk, dict):
                    continue
                btype = blk.get("type")
                if btype == "text" and blk.get("text"):
                    parts.append(str(blk["text"]))
                elif btype == "table" and blk.get("table"):
                    tbl = blk["table"]
                    tbls.append([[str(c) if c is not None else "" for c in (row or [])] for row in tbl if isinstance(row, list)])
                    parts.append("见下表")
            if parts:
                result["body"] = "\n".join(parts)
            if tbls:
                result["tables"] = tbls
            result.pop("blocks", None)
        elif "blocks" in result:
            result.pop("blocks", None)
        children = result.get("children")
        result["children"] = [self.__normalize_node(c) for c in children] if isinstance(children, list) else []
        return result

    def __normalize_content(self, content):
        if not isinstance(content, dict) or not isinstance(content.get("sections"), list):
            return copy.deepcopy(DEFAULT_UTP_CONTENT)
        return {"sections": [self.__normalize_node(s) for s in content["sections"]]}

    def __replace_name(self, node, base, name, skip_titles=None):
        if not name or base == name:
            return
        if skip_titles and str(node.get("title") or "").strip() in skip_titles:
            return
        if node.get("body"):
            node["body"] = node["body"].replace(base, name)
        for tbl in (node.get("tables") or []):
            for row in tbl:
                for i in range(len(row)):
                    if isinstance(row[i], str) and base in row[i]:
                        row[i] = row[i].replace(base, name)
        for c in (node.get("children") or []):
            self.__replace_name(c, base, name, skip_titles=skip_titles)

    def __dhf_file_no(self, prod_id):
        if not prod_id:
            return ""
        row = db.session.execute(
            select(ProdDhf).where(ProdDhf.prod_id == prod_id, ProdDhf.name == DOC_NAME).order_by(ProdDhf.id.asc())
        ).scalars().first()
        if not row:
            row = db.session.execute(
                select(ProdDhf).where(ProdDhf.prod_id == prod_id, ProdDhf.name.like(f"%{DOC_NAME}%")).order_by(ProdDhf.id.asc())
            ).scalars().first()
        return (row.code or "").strip() if row and row.code else ""

    def __plan_cover_date(self, prod_id, version=None):
        """编审批、生效日期、修订修改日期：时间线里「用户测试计划」的最晚一天。
        本版本没有时间线时，用同名产品其它版本的时间线。不用宽泛「用户测试」，避免串到用例/记录/报告。"""
        if not prod_id:
            return ""
        rows = serv_review_util.timeline_rows_with_sibling(prod_id)
        if not rows:
            return ""
        return serv_review_util.review_date(prod_id, ["用户测试计划"], tl_rows=rows, version=version) or ""

    def __fill_revision(self, content, prod_id, version, force=False):
        """文件修订记录首行：修改日期(评审/封面日期)、版本号、首次发布、
        修订人=测试人员（2025.09 前宋月/之后孙家旭）、批准人=研发负责人。打开即覆盖修订人/批准人。"""
        rev_date = self.__plan_cover_date(prod_id, version)
        reviser = "宋月" if serv_review_util._before_202509(rev_date) else "孙家旭"
        # 批准人=研发负责人。本版本没有该角色时，用同一产品名称下其它版本的参与人员。
        approver = (serv_review_util.cover_signer_names(prod_id, DOC_KEY, rev_date).get("批准人") or "") if prod_id else ""
        for s in (content.get("sections") or []):
            if s.get("ref_type") != "revision":
                continue
            tables = s.get("tables") or []
            if tables and isinstance(tables[0], list) and tables[0]:
                t = tables[0]
                cols = len(t[0]) if isinstance(t[0], list) and t[0] else 5
                while len(t) < 6:
                    t.append([""] * cols)
                if serv_review_util.stamp_list_revision_date(t, version, rev_date, prod_id=prod_id, doc_key=DOC_KEY):
                    break
                if len(t) >= 2 and isinstance(t[1], list) and len(t[1]) >= 3:
                    row = t[1]
                    if force or rev_date:
                        row[0] = rev_date or ""
                    elif not str(row[0] or "").strip():
                        row[0] = rev_date
                    if force:
                        if len(row) >= 2: row[1] = str(version or "")
                        if len(row) >= 3: row[2] = "首次发布"
                    else:
                        if version and len(row) >= 2 and not str(row[1] or "").strip(): row[1] = str(version)
                        if len(row) >= 3 and not str(row[2] or "").strip(): row[2] = "首次发布"
                    if len(row) >= 4:
                        row[3] = reviser
                    if len(row) >= 5:
                        row[4] = approver
            break
        return content

    def __fill_test_plan(self, content, prod_id):
        """「测试计划」表：计划完成时间按行「测试活动」匹配时间线。
        用户测试计划负责人：活动含「用户培训」→产品经理；活动/交付物含「用户测试记录」→用户人员；
        其余行→测试人员（与修订人相同：2025.09 前宋月/之后孙家旭）。打开即覆盖。"""
        if not prod_id or not isinstance(content, dict):
            return content
        doc_date = self.__plan_cover_date(prod_id)
        members = db.session.execute(
            select(ProjectMember).where(ProjectMember.prod_id == prod_id)
        ).scalars().all()
        is_user = DOC_KEY == "utp"
        if is_user:
            # 「用例」勿用宽关键字，避免命中其它带「用例+评审」的时间线行（误取更早月份）
            date_kw = {"报告": ["用户测试报告"], "用例": ["用户测试用例"],
                       "计划": ["用户测试计划"], "执行": ["用户测试", "测试"]}

            def role_ok(r):
                return "用户测试" in r
        else:
            date_kw = {"报告": ["软件测试报告"],
                       "用例": ["软件测试用例", "单元测试用例", "集成测试用例", "系统测试用例", "测试用例"],
                       "计划": ["软件测试计划"],
                       "执行": ["软件测试", "系统测试", "集成测试", "单元测试", "测试"]}

            def role_ok(r):
                return ("测试" in r) and ("用户" not in r)
        # 用户人员（角色含「用户测试」）；软件测试计划侧仍为其他测试工程师
        others = []
        pms = []  # 产品经理
        for m in members:
            role = str(m.role or "")
            nm = (m.name or "").strip()
            if not nm:
                continue
            if "产品经理" in role and nm not in pms:
                pms.append(nm)
            if role_ok(role):
                if nm not in ("宋月", "孙家旭") and nm not in others:
                    others.append(nm)

        def to_month(d):
            mt = re.match(r"(\d{4})\D+(\d{1,2})", str(d or ""))
            return "%s.%02d" % (mt.group(1), int(mt.group(2))) if mt else str(d or "")

        def row_kind(activity):
            a = str(activity or "")
            if any(k in a for k in ("测试记录", "回归", "首轮")):
                return "执行"
            for kind in ("报告", "用例", "计划"):
                if kind in a:
                    return kind
            if "测试" in a:
                return "执行"
            return ""

        def row_date(kind):
            if kind in ("报告", "用例", "计划", "执行"):
                return serv_review_util.review_date(prod_id, date_kw[kind]) or doc_date
            return doc_date

        def is_user_record_row(act, deliverable=""):
            return "用户测试记录" in f"{act or ''} {deliverable or ''}"

        def owners(dt, kind, act="", deliverable=""):
            if is_user:
                if "用户培训" in str(act or ""):
                    names = list(pms)
                elif is_user_record_row(act, deliverable):
                    names = list(others)
                else:
                    # 与修订人同一逻辑：测试人员日期规则
                    names = ["宋月" if serv_review_util._before_202509(dt or doc_date) else "孙家旭"]
            else:
                primary = "宋月" if serv_review_util._before_202509(dt or doc_date) else "孙家旭"
                # 评审计划/测试用例/测试报告：只写主测试工程师一人；其余(执行)：主+其他测试人员
                if kind in ("报告", "用例", "计划"):
                    names = [primary]
                else:
                    names = [primary] + [n for n in others if n != primary]
            return "、".join([n for n in names if n])

        def walk(node):
            for tbl in (node.get("tables") or []):
                if not tbl or not isinstance(tbl[0], list):
                    continue
                header = [str(c or "") for c in tbl[0]]
                hj = " ".join(header)
                if "计划完成时间" in hj and "负责人" in hj:
                    ci_time = next((i for i, h in enumerate(header) if "计划完成时间" in h), 0)
                    ci_own = next((i for i, h in enumerate(header) if "负责人" in h), 1)
                    ci_act = next((i for i, h in enumerate(header) if "活动" in h), 2)
                    ci_del = next((i for i, h in enumerate(header) if "交付" in h), -1)
                    for r in tbl[1:]:
                        if not isinstance(r, list):
                            continue
                        act = r[ci_act] if ci_act < len(r) else ""
                        deliverable = r[ci_del] if ci_del >= 0 and ci_del < len(r) else ""
                        kind = row_kind(act)
                        dt = to_month(row_date(kind))
                        if ci_time < len(r):
                            r[ci_time] = dt
                        if ci_own < len(r):
                            r[ci_own] = owners(dt, kind, act, deliverable)
            for ch in (node.get("children") or []):
                walk(ch)

        for s in (content.get("sections") or []):
            walk(s)
        return content

    def __fill_roles(self, content, prod_id, force=False):
        """「角色和职责」表：按「职责」与测试计划负责人同一规则覆盖「角色」「人员」。
        用户培训→产品经理；用户测试记录→模拟用户；其余→测试人员。
        force=False：匹配不到则保留原值；force=True：强制覆盖，无人员则置空。"""
        if not isinstance(content, dict):
            return content
        doc_date = self.__plan_cover_date(prod_id)
        members = db.session.execute(
            select(ProjectMember).where(ProjectMember.prod_id == prod_id)
        ).scalars().all() if prod_id else []
        primary = "宋月" if serv_review_util._before_202509(doc_date) else "孙家旭"
        pms = []
        utesters = []
        for m in members:
            role = str(m.role or "")
            nm = (m.name or "").strip()
            if not nm:
                continue
            if "产品经理" in role and nm not in pms:
                pms.append(nm)
            if "用户测试" in role and nm not in utesters:
                utesters.append(nm)

        def role_and_names(duty):
            a = str(duty or "")
            if "用户培训" in a:
                return "产品经理", "、".join(pms) if pms else ""
            if "用户测试记录" in a:
                return "模拟用户", "、".join(utesters) if utesters else ""
            # 计划/用例/报告及评审等 → 测试人员（与测试计划负责人一致）
            return "测试人员", primary

        def walk(node):
            for tbl in (node.get("tables") or []):
                if not tbl or not isinstance(tbl[0], list):
                    continue
                header = [str(c or "") for c in tbl[0]]
                hj = " ".join(header)
                if "人员" in hj and "角色" in hj and "职责" in hj:
                    ci_p = next((i for i, h in enumerate(header) if "人员" in h), 0)
                    ci_r = next((i for i, h in enumerate(header) if "角色" in h), 1)
                    ci_d = next((i for i, h in enumerate(header) if "职责" in h), 2)
                    for r in tbl[1:]:
                        if not isinstance(r, list):
                            continue
                        duty = r[ci_d] if ci_d < len(r) else ""
                        role_label, nm = role_and_names(duty)
                        if force:
                            if ci_r < len(r):
                                r[ci_r] = role_label
                            if ci_p < len(r):
                                r[ci_p] = nm or ""
                        else:
                            if nm and ci_p < len(r):
                                r[ci_p] = nm
                            if role_label and ci_r < len(r):
                                r[ci_r] = role_label
            for ch in (node.get("children") or []):
                walk(ch)

        for s in (content.get("sections") or []):
            walk(s)
        return content

    def __fill_runtime_env(self, content, prod_id, force=False):
        """运行环境表从产品 prod_runtime_env 整表覆盖（含增删行列后的结构）。"""
        return apply_runtime_to_pdp_content(content, prod_id)

    def __clean_plan_captions(self, content):
        """「测试计划」表名去掉前缀版本号残片（如 .0.0测试计划如下：→ 测试计划如下：）。"""
        if not isinstance(content, dict):
            return content

        def walk(node):
            caps = node.get("table_captions")
            if isinstance(caps, list) and caps:
                fixed = []
                for c in caps:
                    s = str(c or "")
                    if "测试计划如下" in s:
                        s = "测试计划如下："
                    fixed.append(s)
                node["table_captions"] = fixed
            for ch in (node.get("children") or []):
                walk(ch)

        for s in (content.get("sections") or []):
            walk(s)
        return content

    @staticmethod
    def __strip_title(title):
        return re.sub(r"^\s*\d+(?:\.\d+)*[、.\s]*", "", str(title or "")).strip()

    def __fill_desc(self, content, product):
        """「描述」章：产品名称/完整版本自动写入；有产品型号(type_code)则写入，无则隐藏该行；保留其后说明正文。"""
        if not product:
            return
        name = (product.name or "").strip()
        type_code = (getattr(product, "type_code", None) or "").strip()
        full_version = (product.full_version or "").strip()
        if not (name or type_code or full_version):
            return

        def walk(node):
            if self.__strip_title(node.get("title")) == "描述":
                body = str(node.get("body") or "")
                rest = []
                for ln in body.split("\n"):
                    s = ln.strip()
                    if s.startswith("产品名称：") or s.startswith("产品型号：") or s.startswith("完整版本："):
                        continue
                    rest.append(ln)
                while rest and not str(rest[0]).strip():
                    rest.pop(0)
                header = []
                if name:
                    header.append(f"产品名称：{name}")
                if type_code:
                    header.append(f"产品型号：{type_code}")
                if full_version:
                    header.append(f"完整版本：{full_version}")
                node["body"] = "\n".join(header + rest) if header else "\n".join(rest)
            for c in (node.get("children") or []):
                walk(c)

        for s in (content.get("sections") or []):
            walk(s)

    def __autofill(self, content, prod_id, product=None, version="", force=False):
        if not isinstance(content, dict):
            return content
        if product and product.name:
            for s in (content.get("sections") or []):
                self.__replace_name(s, BASE_NAME, product.name)
            # 更新完整版本号（「描述」章由 __fill_desc 重建头部）
            new_full_version = (product.full_version or "").strip()
            if new_full_version:
                def _update_version(node):
                    if self.__strip_title(node.get("title")) == "描述":
                        pass
                    elif node.get("body") and "完整版本" in str(node.get("body")):
                        node["body"] = re.sub("完整版本：[^\n]*", f"完整版本：{new_full_version}", str(node["body"]))
                    for c in (node.get("children") or []):
                        _update_version(c)
                for s in (content.get("sections") or []):
                    _update_version(s)
        self.__fill_desc(content, product)
        self.__clean_plan_captions(content)
        self.__fill_test_plan(content, prod_id)
        self.__fill_roles(content, prod_id, force=force)
        self.__fill_runtime_env(content, prod_id, force=force)
        self.__fill_revision(content, prod_id, version, force=force)
        serv_review_util.ensure_review(
            content, DOC_KEY,
            serv_review_util.review_date(prod_id, serv_review_util.REVIEW_DEFS[DOC_KEY]["name_keywords"], version=version) if prod_id else "",
            prod_id,
            force_date=bool(serv_review_util.raised_doc_version(version)),
        )
        cover_date = self.__plan_cover_date(prod_id, version)
        serv_review_util.fill_cover_dates(
            content, cover_date,
            force=force or bool(serv_review_util.raised_doc_version(version) and cover_date) or serv_review_util.cover_date_overwrite(version, cover_date),
        )
        serv_review_util.fill_cover_signers(content, serv_review_util.cover_signers(prod_id, DOC_KEY) if prod_id else {}, force=force)
        return content

    def __to_obj(self, row: UtpDoc, product: Product = None):
        obj = UtpDocObj(**row.dict())
        obj.content = self.__autofill(self.__normalize_content(obj.content), row.product_id, product, row.version)
        if not (obj.file_no or "").strip():
            obj.file_no = self.__dhf_file_no(row.product_id)
        if product:
            obj.product_name = product.name
            obj.product_version = product.full_version
            obj.product_full_version = product.full_version
            obj.product_type_code = product.type_code
        obj.file_no = sync_file_no_version(obj.file_no or "", getattr(row, "version", None) or "") or obj.file_no
        return obj

    async def add_utp_doc(self, form: UtpDocForm):
        try:
            sql = select(func.count(UtpDoc.id)).where(UtpDoc.product_id == form.product_id, UtpDoc.version == form.version)
            if db.session.execute(sql).scalar() > 0:
                return Resp.resp_err(msg=ts("msg_obj_exist"))
            row = UtpDoc(**form.dict(exclude_none=True))
            row.id = None
            row.content = self.__normalize_content(row.content)
            db.session.add(row)
            db.session.commit()
            return Resp.resp_ok(data=UtpDocForm(id=row.id))
        except Exception:
            logger.exception("")
            db.session.rollback()
        return Resp.resp_err(msg=ts(msg_err_db))

    async def duplicate_utp_doc(self, id: int, product_id: int = None):
        try:
            fromdoc: UtpDoc = db.session.execute(select(UtpDoc).where(UtpDoc.id == id)).scalars().first()
            if not fromdoc:
                return Resp.resp_err(msg=ts("msg_obj_null"))
            target_pid = product_id or fromdoc.product_id
            all_versions = db.session.execute(select(UtpDoc.version).where(UtpDoc.product_id == target_pid)).scalars().all()
            existing_set = {v for v in all_versions if v}
            if target_pid == fromdoc.product_id:
                version = new_version(fromdoc.version)
            else:
                def _seq(v):
                    m = re.search(r"(\d+)(?!.*\d)", v or "")
                    return int(m.group(1)) if m else -1
                valid = [v for v in all_versions if v]
                version = new_version(max(valid, key=_seq)) if valid else fromdoc.version
            while version in existing_set:
                version = new_version(version)
            newdoc = UtpDoc(
                product_id=target_pid, version=version,
                file_no=sync_file_no_version((fromdoc.file_no or "").strip() or self.__dhf_file_no(target_pid), version) or None,
                change_log=fromdoc.change_log,
                content=copy.deepcopy(self.__normalize_content(fromdoc.content)),
            )
            db.session.add(newdoc)
            serv_review_util.add_raised_copy_dhf(
                target_pid, fromdoc.product_id, getattr(fromdoc, "file_no", None), version, DOC_NAME,
            )
            db.session.commit()
            if target_pid != fromdoc.product_id:
                await self.rebind_product(newdoc.id, target_pid)
            return Resp.resp_ok(data=UtpDocForm(id=newdoc.id))
        except Exception:
            logger.exception("")
            db.session.rollback()
        return Resp.resp_err(msg=ts(msg_err_db))

    async def update_utp_doc(self, form: UtpDocForm):
        try:
            row: UtpDoc = db.session.execute(select(UtpDoc).where(UtpDoc.id == form.id)).scalars().first()
            if not row:
                return Resp.resp_err(msg=ts("msg_obj_null"))
            for key, value in form.dict(exclude_none=True).items():
                if key == "id":
                    continue
                if key == "content":
                    value = self.__normalize_content(value)
                setattr(row, key, value)
            if getattr(form, "version", None) and hasattr(row, "file_no"):
                row.file_no = sync_file_no_version(row.file_no or "", form.version or row.version) or row.file_no
            db.session.commit()
            return Resp.resp_ok()
        except Exception:
            logger.exception("")
            db.session.rollback()
        return Resp.resp_err(msg=ts(msg_err_db))

    async def rebind_product(self, id: int, product_id: int):
        """切换产品：更新 product_id 并强制用新产品信息重新获取封面/修订/产品名后保存，返回新 obj。"""
        try:
            row: UtpDoc = db.session.execute(select(UtpDoc).where(UtpDoc.id == id)).scalars().first()
            if not row:
                return Resp.resp_err(msg=ts("msg_obj_null"))
            old_product: Product = db.session.execute(select(Product).where(Product.id == row.product_id)).scalars().first() if row.product_id else None
            product: Product = db.session.execute(select(Product).where(Product.id == product_id)).scalars().first()
            if not product:
                return Resp.resp_err(msg=ts("msg_obj_null"))
            db.session.execute(delete(UtpDoc).where(UtpDoc.product_id == product_id, UtpDoc.version == row.version, UtpDoc.id != id))
            content = self.__normalize_content(row.content)
            # 重置含产品名/完整版本的固定章节为模板原值（恢复基准名+原版本号，避免被旧产品污染）
            tpl = copy.deepcopy(DEFAULT_UTP_CONTENT) if isinstance(DEFAULT_UTP_CONTENT, dict) else {"sections": []}
            tpl_map = {}
            def collect_tpl(node):
                key = str(node.get("title") or "").strip()
                body = str(node.get("body") or "")
                caps = node.get("table_captions") or []
                if BASE_NAME in body or "完整版本" in body or any(BASE_NAME in str(r) for tbl in (node.get("tables") or []) for r in tbl) or any("1.1.0.1" in str(c) for c in caps):
                    tpl_map[key] = {"body": node.get("body", ""), "tables": copy.deepcopy(node.get("tables", [])), "table_captions": copy.deepcopy(caps)}
                for c in (node.get("children") or []):
                    collect_tpl(c)
            for s in (tpl.get("sections") or []):
                collect_tpl(s)
            def reset_fixed(node):
                key = str(node.get("title") or "").strip()
                if key in tpl_map:
                    node["body"] = tpl_map[key]["body"]
                    node["tables"] = copy.deepcopy(tpl_map[key]["tables"])
                    if tpl_map[key].get("table_captions"):
                        node["table_captions"] = copy.deepcopy(tpl_map[key]["table_captions"])
                for c in (node.get("children") or []):
                    reset_fixed(c)
            for s in (content.get("sections") or []):
                reset_fixed(s)
            # 完整版本号更新为新产品版本（body + table_captions；「测试计划如下」不含版本号）
            new_full_version = (product.full_version or "").strip()
            def update_version(node):
                if node.get("body") and "完整版本" in str(node.get("body")):
                    node["body"] = re.sub("完整版本：[^\n]*", f"完整版本：{new_full_version}", str(node["body"]))
                caps = node.get("table_captions") or []
                if caps:
                    next_caps = []
                    for c in caps:
                        s = str(c or "")
                        if "测试计划如下" in s:
                            next_caps.append("测试计划如下：")
                        elif new_full_version and re.search(r"\d+\.\d+\.\d+\.\d+", s):
                            next_caps.append(re.sub(r"\d+\.\d+\.\d+\.\d+", new_full_version, s))
                        else:
                            next_caps.append(c)
                    node["table_captions"] = next_caps
                for c in (node.get("children") or []):
                    update_version(c)
            for s in (content.get("sections") or []):
                update_version(s)
            row.product_id = product_id
            content = self.__autofill(content, product_id, product, row.version, force=True)
            row.content = content
            db.session.commit()
            return Resp.resp_ok(data=self.__to_obj(row, product))
        except Exception:
            logger.exception("")
            db.session.rollback()
        return Resp.resp_err(msg=ts(msg_err_db))

    async def delete_utp_doc(self, id: int):
        db.session.execute(delete(UtpDoc).where(UtpDoc.id == id))
        db.session.commit()
        return Resp.resp_ok()

    async def get_utp_doc(self, id: int):
        sql = select(UtpDoc, Product).join(Product, UtpDoc.product_id == Product.id).where(UtpDoc.id == id)
        row = db.session.execute(sql).first()
        if not row:
            return Resp.resp_err(msg=ts("msg_obj_null"))
        doc, product = row
        return Resp.resp_ok(data=self.__to_obj(doc, product))

    async def list_utp_doc(self, op_user: UserObj = None, product_id: int = 0, version: str = None, page_index: int = 0, page_size: int = 10):
        wheres = []
        if product_id:
            wheres.append(UtpDoc.product_id == product_id)
        if version:
            wheres.append(UtpDoc.version.like(f"%{version}%"))
        scope = product_scope_filter(op_user)
        if scope is not None:
            wheres.append(scope)
        sql_total = select(func.count(UtpDoc.id)).join(Product, UtpDoc.product_id == Product.id).where(*wheres)
        total = db.session.execute(sql_total).scalar() or 0
        sql = (select(UtpDoc, Product).join(Product, UtpDoc.product_id == Product.id).where(*wheres)
               .order_by(UtpDoc.id.desc()).offset(page_index * page_size).limit(page_size))
        rows: List[UtpDocObj] = [self.__to_obj(doc, product) for doc, product in db.session.execute(sql).all()]
        return Resp.resp_ok(data=Page(total=total, rows=rows, page_index=page_index, page_size=page_size))

    # ---------------- 导出 Word（PDP 风格章节树） ----------------
    async def export_utp_doc(self, output, id: int):
        resp = await self.get_utp_doc(id)
        obj = resp.data
        if obj is None:
            Document().save(output)
            output.seek(0)
            return
        c = self.__normalize_content(obj.content)
        sections = c.get("sections") or []
        document = Document()
        section = document.sections[0]
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        update_fields = OxmlElement("w:updateFields")
        update_fields.set(qn("w:val"), "true")
        document.settings.element.append(update_fields)
        header_para = section.header.add_paragraph()
        header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        docx_util.fonted_txt(header_para, obj.file_no or "")
        docx_util.add_page_number_footer(section, obj.file_no or "")

        def add_blank_lines(n):
            for _ in range(max(0, int(n or 0))):
                document.add_paragraph("")

        def write_center_title(text, size=22.0, bold=False):
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            docx_util.fonted_txt(p, str(text or ""), font_size=size, bold=bold)

        def add_text(text):
            docx_util.save_txt2docx(str(text or ""), document)

        def set_cell(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
            s = str(text or "")
            if s.startswith("data:image"):
                try:
                    b64 = s.split(",", 1)[1] if "," in s else ""
                    cell.text = ""
                    para = cell.paragraphs[0]
                    para.alignment = align
                    para.add_run().add_picture(BytesIO(base64.b64decode(b64)), height=Pt(33))
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    return
                except Exception:
                    pass
            cell.text = ""
            for i, line in enumerate(s.split("\n")):
                para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
                para.alignment = align
                para.paragraph_format.line_spacing = 1.3
                docx_util.fonted_txt(para, line, font_size=10.5, bold=bold)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER if align == WD_ALIGN_PARAGRAPH.CENTER else WD_CELL_VERTICAL_ALIGNMENT.TOP

        def _set_fixed_widths(table, widths_dxa):
            table.autofit = False
            table.allow_autofit = False
            tbl_pr = table._tbl.tblPr
            layout = tbl_pr.find(qn("w:tblLayout"))
            if layout is None:
                layout = OxmlElement("w:tblLayout")
                tbl_pr.append(layout)
            layout.set(qn("w:type"), "fixed")
            grid_el = table._tbl.find(qn("w:tblGrid"))
            if grid_el is not None:
                for gc in list(grid_el):
                    grid_el.remove(gc)
                for w in widths_dxa:
                    gc = OxmlElement("w:gridCol")
                    gc.set(qn("w:w"), str(w))
                    grid_el.append(gc)
            # 逐格写 tcW，确保固定列宽生效
            for row in table.rows:
                for i, w in enumerate(widths_dxa):
                    if i < len(row.cells):
                        tcpr = row.cells[i]._tc.get_or_add_tcPr()
                        tcw = tcpr.find(qn("w:tcW"))
                        if tcw is None:
                            tcw = OxmlElement("w:tcW")
                            tcpr.append(tcw)
                        tcw.set(qn("w:w"), str(w))
                        tcw.set(qn("w:type"), "dxa")

        def add_grid(grid):
            grid = [row for row in (grid or []) if isinstance(row, list)]
            cols = max((len(row) for row in grid), default=0)
            if cols <= 0:
                return
            table = document.add_table(rows=0, cols=cols)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True
            for r_idx, row in enumerate(grid):
                cells = table.add_row().cells
                for c_idx in range(cols):
                    set_cell(cells[c_idx], row[c_idx] if c_idx < len(row) else "", bold=(r_idx == 0))
            # SCI 清单表：固定列宽，SCI名字/类型不换行、存储地点较宽
            first = str(grid[0][0]).strip() if grid and grid[0] else ""
            if first == "SCI名字":
                if cols == 5:
                    _set_fixed_widths(table, [1700, 1100, 2000, 1500, 3000])
                elif cols == 4:
                    _set_fixed_widths(table, [1700, 2200, 1500, 3600])
            elif first == "类别" and cols >= 3:
                # 配置项存储地址表：类别列收窄、存储路径较宽；首列连续相同「类别」纵向合并；文字左对齐
                if cols == 3:
                    _set_fixed_widths(table, [1300, 1300, 5000])
                rows = table.rows
                n = len(rows)
                r = 1
                while r < n:
                    val = str((grid[r][0] if r < len(grid) and grid[r] else "") or "").strip()
                    if not val:
                        r += 1
                        continue
                    r2 = r
                    while r2 + 1 < n and str((grid[r2 + 1][0] if grid[r2 + 1] else "") or "").strip() == val:
                        r2 += 1
                    if r2 > r:
                        merged = rows[r].cells[0].merge(rows[r2].cells[0])
                        set_cell(merged, val, align=WD_ALIGN_PARAGRAPH.LEFT)
                        merged.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    r = r2 + 1
            elif first in ("测试目标", "测试方法和技术") and cols == 2:
                _set_fixed_widths(table, [2200, 7100])
            elif first == "配置":
                if cols == 2:
                    _set_fixed_widths(table, [1900, 7400])
                elif cols == 3:
                    _set_fixed_widths(table, [1900, 3700, 3700])
            document.add_paragraph()

        def add_cover_grid(grid):
            grid = [row for row in (grid or []) if isinstance(row, list)]
            cols = max((len(row) for row in grid), default=0)
            if cols <= 0:
                return
            table = document.add_table(rows=0, cols=cols)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True
            for row in grid:
                cells = table.add_row().cells
                for c_idx in range(cols):
                    text = row[c_idx] if c_idx < len(row) else ""
                    set_cell(cells[c_idx], text, bold=(c_idx % 2 == 0), align=WD_ALIGN_PARAGRAPH.CENTER)
                if (str(row[0]).strip() if row else "") == "生效日期" and cols > 2:
                    merged = cells[1]
                    for c_idx in range(2, cols):
                        merged = merged.merge(cells[c_idx])
                    set_cell(merged, row[1] if len(row) > 1 else "", align=WD_ALIGN_PARAGRAPH.CENTER)
            document.add_paragraph()

        def strip_num(title):
            return re.sub(r"^\s*\d+(?:\.\d+)*[\.、\s]*", "", str(title or "")).strip()

        def add_body_heading(title, level):
            size = {1: 16.0, 2: 14.0, 3: 12.0}.get(level, 11.0)
            p = document.add_heading("", level=max(1, min(level, 9)))
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.left_indent = Pt(0)
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            docx_util.fonted_txt(p, title, font_size=size, bold=True)

        def add_image(image_url):
            s = str(image_url or "")
            if not s.startswith("data:image"):
                return
            try:
                b64 = s.split(",", 1)[1] if "," in s else ""
                raw = base64.b64decode(b64)
                p = document.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(BytesIO(raw), width=Inches(4.5))
                document.add_paragraph("")
            except Exception:
                logger.exception("导出章节图片失败")

        def clean_export_text(text):
            """去掉编辑页静默占位，避免 Word 露出 {{TBL}}/{{TXT}} 等。"""
            lines = []
            for ln in str(text or "").split("\n"):
                s = ln.strip()
                if s in ("{{TBL}}", "{{IMG}}", "{{EMPTY}}", "{{TXT}}"):
                    continue
                lines.append(ln)
            return "\n".join(lines).strip()

        def render_body_section(node, level, number=""):
            name = strip_num(node.get("title"))
            heading = f"{number} {name}".strip() if number else name
            add_body_heading(heading, level=max(1, min(level, 9)))
            if node.get("ref_type") == "review":
                for t_idx, table in enumerate(node.get("tables") or []):
                    serv_review_util.render_review_grid(document, table, set_cell, merge_col0=(t_idx == 0), merge_full=True)
            else:
                body_text = str(node.get("body") or "")
                tables = node.get("tables") or []
                caps = node.get("table_captions") or []
                images = [x for x in (node.get("images") or []) if isinstance(x, str) and x]

                def _emit_table(t_idx):
                    if t_idx < 0 or t_idx >= len(tables):
                        return
                    cap = caps[t_idx] if t_idx < len(caps) else ""
                    if str(cap or "").strip():
                        add_text(cap)
                    add_grid(tables[t_idx])

                # 与编辑页交错一致：{{TBL}}/{{IMG}}/见下表/表N；静默占位不写入正文
                has_tbl_ph = "{{TBL}}" in body_text
                has_img_ph = "{{IMG}}" in body_text
                if (has_tbl_ph or has_img_ph or ("见下表" in body_text)
                        or re.search(r"见表\s*\d", body_text) or re.search(r"(?m)^表\s*\d", body_text)
                        or tables or images):
                    lines = body_text.split("\n")
                    buf = []
                    tbl_idx = 0
                    img_idx = 0

                    def flush_text():
                        t = clean_export_text("\n".join(buf))
                        if t:
                            add_text(t)
                        buf.clear()

                    for ln in lines:
                        s = ln.strip()
                        if (s == "{{TBL}}" or s in ("见下表", "如下表", "下表") or re.match(r"^表\s*\d", s)) and tbl_idx < len(tables):
                            flush_text()
                            _emit_table(tbl_idx)
                            tbl_idx += 1
                        elif re.search(r"见表\s*\d", s) and tbl_idx < len(tables):
                            rest = re.sub(r"见表\s*\d+", "", ln).strip()
                            if rest:
                                buf.append(rest)
                            flush_text()
                            _emit_table(tbl_idx)
                            tbl_idx += 1
                        elif (s == "{{IMG}}" or s in ("见下图", "如下图", "下图") or re.match(r"^图\s*\d", s)) and img_idx < len(images):
                            flush_text()
                            add_image(images[img_idx])
                            img_idx += 1
                        elif (
                            not has_tbl_ph
                            and re.match(r"^\d+[）)]", s)
                            and tbl_idx < len(tables)
                        ):
                            buf.append(ln)
                            flush_text()
                            _emit_table(tbl_idx)
                            tbl_idx += 1
                        elif s in ("{{TBL}}", "{{IMG}}", "{{EMPTY}}", "{{TXT}}"):
                            continue
                        else:
                            buf.append(ln)
                    flush_text()
                    for i in range(tbl_idx, len(tables)):
                        _emit_table(i)
                    for i in range(img_idx, len(images)):
                        add_image(images[i])
                else:
                    cleaned = clean_export_text(body_text)
                    if cleaned:
                        add_text(cleaned)
                if str(node.get("body_tail") or "").strip():
                    add_text(clean_export_text(node.get("body_tail")))
            idx = 0
            for child in (node.get("children") or []):
                idx += 1
                child_num = f"{number}.{idx}" if number else f"{idx}"
                render_body_section(child, level + 1, child_num)

        cover = next((s for s in sections if s.get("ref_type") == "cover"), None)
        revision = next((s for s in sections if s.get("ref_type") == "revision"), None)
        body = [s for s in sections if s.get("ref_type") not in ("cover", "revision")]

        add_blank_lines(6)
        write_center_title((strip_num(cover.get("title")) if cover else "") or DOC_NAME, size=22.0, bold=True)
        add_blank_lines(4)
        if cover:
            for table in (cover.get("tables") or []):
                add_cover_grid(table)

        document.add_page_break()
        write_center_title("文件修订记录", size=14.0, bold=True)
        add_blank_lines(2)
        if revision:
            for table in (revision.get("tables") or []):
                add_grid(table)

        document.add_page_break()
        write_center_title("目录", size=16.0, bold=True)
        docx_util.insert_toc_field(document)

        document.add_page_break()
        seq = 0
        for node in body:
            if node.get("ref_type") == "review":
                render_body_section(node, 1, "")
            else:
                seq += 1
                render_body_section(node, 1, str(seq))

        docx_util.fill_toc_cache(document)
        document.save(output)
        output.seek(0)
