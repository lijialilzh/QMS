#!/usr/bin/env python
# encoding: utf-8

# 现场测试记录服务层（测试文件 VV-006）。
# 默认内容取自 src-res/ftr_record_default_content.json（一张大表格）。

import copy
import io
import json
from collections import Counter
import logging
import os
import re
import urllib.parse
from datetime import datetime
from typing import Any, List

from sqlalchemy import delete, func, select
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT

from ..model.product import Product
from ..model.ftr_doc import FtrDoc
from ..model.ftr_record_doc import FtrRecordDoc
from ..model.prod_dhf import ProdDhf
from ..model.project_member import ProjectMember
from ..obj import Page, Resp
from ..obj.tobj_role import Roles
from ..obj.vobj_user import UserObj
from ..obj.tobj_ftr_record_doc import FtrRecordDocForm
from ..obj.vobj_ftr_record_doc import FtrRecordDocObj
from ..utils.i18n import ts
from ..utils.sql_ctx import db
from . import msg_err_db
from .serv_utils import new_version, sync_file_no_version
from .serv_utils import docx_util
from . import serv_review_util
from .serv_product import product_scope_filter

logger = logging.getLogger(__name__)

DOC_NAME = "现场测试记录"
DOC_KEY = "ftr_record"
BASE_NAME = "InferOperate Suite"
_PURPOSE_ITEMS = ("成品检验", "现场测试")
_EMPTY_RESULT = "□通过   □不通过"

_DEFAULT_FILE = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "src-res", "ftr_record_default_content.json"
)
try:
    with open(_DEFAULT_FILE, encoding="utf-8") as _f:
        DEFAULT_CONTENT = json.load(_f)
except Exception:
    DEFAULT_CONTENT = {"sections": []}


def _default_case_keys():
    """模板里的功能模块+功能，同步时丢掉，避免旧样例和规程叠在一起。"""
    keys = set()
    try:
        grid = (((DEFAULT_CONTENT.get("sections") or [{}])[0].get("tables") or [[]])[0])
    except Exception:
        return keys
    start = 7
    for i, row in enumerate(grid or []):
        if isinstance(row, list) and row and str(row[0] or "").strip() == "功能模块":
            start = i + 1
            break
    for row in (grid or [])[start:]:
        if isinstance(row, list):
            keys.add((str(row[0] or "").strip(), str(row[1] or "").strip()))
    return keys


class Server(object):

    def __normalize_content(self, content):
        if not isinstance(content, dict) or not isinstance(content.get("sections"), list):
            return copy.deepcopy(DEFAULT_CONTENT)
        result = {"sections": []}
        for s in content["sections"]:
            node = dict(s)
            node["title"] = str(node.get("title") or "")
            node["body"] = str(node.get("body") or "")
            tables = node.get("tables")
            if not isinstance(tables, list):
                tables = []
            norm = []
            for t in tables:
                if isinstance(t, list):
                    norm.append([[str(c) if c is not None else "" for c in (r or [])] for r in t if isinstance(r, list)])
            node["tables"] = norm
            node["children"] = []
            result["sections"].append(node)
        return result

    def __replace_name(self, node, base, name):
        if not name or base == name:
            return
        if node.get("body"):
            node["body"] = node["body"].replace(base, name)
        for tbl in (node.get("tables") or []):
            for row in tbl:
                for i in range(len(row)):
                    if isinstance(row[i], str) and base in row[i]:
                        row[i] = row[i].replace(base, name)

    def __dhf_file_no(self, prod_id):
        if not prod_id:
            return ""
        row = db.session.execute(
            select(ProdDhf).where(ProdDhf.prod_id == prod_id, ProdDhf.name == DOC_NAME).order_by(ProdDhf.id.asc())
        ).scalars().first()
        if not row:
            row = db.session.execute(
                select(ProdDhf).where(ProdDhf.prod_id == prod_id, ProdDhf.name.like(f"%{DOC_NAME}%")).order_by(ProdDhf.id.asc())
            ).scalars().first()
        return (row.code or "").strip() if row and row.code else ""

    @staticmethod
    def __plain_title(title):
        return re.sub(r"^\s*(第[一二三四五六七八九十0-9]+章\s*)?(?:\d+(?:\.\d+)*[、.\s]*)?", "", str(title or "")).strip()

    @staticmethod
    def __case_titles(tables):
        titles = []
        for grid in tables or []:
            if not isinstance(grid, list):
                continue
            for row in grid:
                if isinstance(row, list) and row and str(row[0] or "").strip() == "用例标题":
                    title = str(row[1] if len(row) > 1 else "").strip()
                    if title:
                        titles.append(title)
                    break
        return titles

    def __ftr_cases(self, prod_id, version=""):
        """同产品《现场测试规程》第2章：功能模块=「功能测试用例」下的二级章节名，功能=该节及其下级的用例标题。"""
        if not prod_id:
            return []
        rows = db.session.execute(
            select(FtrDoc).where(FtrDoc.product_id == prod_id).order_by(FtrDoc.id.desc())
        ).scalars().all()
        if not rows:
            return []
        picked = next((r for r in rows if version and str(r.version or "") == str(version)), None) or rows[0]
        content = picked.content if isinstance(picked.content, dict) else {}
        sections = content.get("sections") or []
        roots = [s for s in sections if isinstance(s, dict) and "功能测试用例" in self.__plain_title(s.get("title"))]
        found = []

        def collect(node, module):
            for title in self.__case_titles(node.get("tables")):
                found.append((module, title))
            for child in (node.get("children") or []):
                if isinstance(child, dict):
                    collect(child, module)

        for root in roots:
            for child in (root.get("children") or []):
                if not isinstance(child, dict):
                    continue
                module = self.__plain_title(child.get("title"))
                if not module or "功能测试用例" in module:
                    continue
                collect(child, module)
        return found

    @staticmethod
    def __record_grid(content):
        sections = content.get("sections") or []
        if not sections or not isinstance(sections[0], dict):
            return None, None
        tables = sections[0].get("tables") or []
        if not tables or not isinstance(tables[0], list):
            return None, None
        return sections[0], tables[0]

    @staticmethod
    def __data_start(grid):
        for i, row in enumerate(grid or []):
            if isinstance(row, list) and row and str(row[0] or "").strip() == "功能模块":
                return i + 1
        return 7

    @staticmethod
    def __format_purpose(text):
        raw = str(text or "")
        parts = []
        for name in _PURPOSE_ITEMS:
            checked = bool(re.search(r"[■☑√✓]\s*" + re.escape(name), raw))
            parts.append(("■" if checked else "□") + name)
        return "   ".join(parts)

    @staticmethod
    def __is_template_case_rows(rows):
        """测试项目仍是模板样例（含重复行）时才允许按规程重建。"""
        keys = []
        for row in rows or []:
            if not isinstance(row, list):
                continue
            module = str(row[0] if row else "").strip()
            func = str(row[1] if len(row) > 1 else "").strip()
            if not module and not func:
                continue
            keys.append((module, func))
        default_keys = _default_case_keys()
        if not keys or not default_keys:
            return False
        return Counter(keys) == Counter(default_keys)

    def __sync_cases(self, content, prod_id, version=""):
        """有现场测试规程时，按规程重建功能行；对得上的测试结果和备注保留。没有规程则不改行。"""
        section, grid = self.__record_grid(content)
        if section is None or not isinstance(grid, list) or len(grid) < 5:
            return
        purpose = self.__format_purpose(grid[4][1] if len(grid[4]) > 1 else "")
        while len(grid[4]) < 5:
            grid[4].append("")
        for i in range(1, 5):
            grid[4][i] = purpose
        cases = self.__ftr_cases(prod_id, version)
        if not cases:
            return
        start = self.__data_start(grid)
        old = [r for r in grid[start:] if isinstance(r, list)]
        # 用户已经增删改过的列表按保存顺序保留，不再按规程把改过的行挪到表尾
        if not self.__is_template_case_rows(old):
            return
        buckets = {}
        by_func = {}
        for row in old:
            key = (str(row[0] or "").strip(), str(row[1] or "").strip())
            buckets.setdefault(key, []).append(row)
            by_func.setdefault(key[1], []).append(row)
        used = set()
        new_rows = []
        for module, func in cases:
            prev = buckets.get((module, func)) or []
            row = prev.pop(0) if prev else None
            if row is None:
                alt = by_func.get(func) or []
                while alt and id(alt[0]) in used:
                    alt.pop(0)
                row = alt.pop(0) if alt else None
            if row is not None:
                used.add(id(row))
            result = str((row[2] if row and len(row) > 2 else "") or _EMPTY_RESULT)
            remark = str(row[4] if row and len(row) > 4 else "")
            new_rows.append([module, func, result, result, remark])
        # 规程里没有、模板里也没有的行是手增的，留在表尾
        default_keys = _default_case_keys()
        for row in old:
            if id(row) in used:
                continue
            key = (str(row[0] or "").strip(), str(row[1] or "").strip())
            if key in default_keys or not key[1]:
                continue
            result = str((row[2] if len(row) > 2 else "") or _EMPTY_RESULT)
            remark = str(row[4] if len(row) > 4 else "")
            new_rows.append([key[0], key[1], result, result, remark])
        section["tables"][0] = grid[:start] + new_rows

    def __autofill(self, content, prod_id, product=None, version="", force=False):
        if not isinstance(content, dict):
            return content
        if product and product.name:
            for s in (content.get("sections") or []):
                self.__replace_name(s, BASE_NAME, product.name)
        # 完整版本号更新
        if product:
            new_fv = (product.full_version or "").strip()
            for s in (content.get("sections") or []):
                for tbl in (s.get("tables") or []):
                    for r in tbl:
                        for i in range(len(r)):
                            if isinstance(r[i], str) and re.search(r"\d+\.\d+\.\d+\.\d+", r[i]):
                                r[i] = re.sub(r"\d+\.\d+\.\d+\.\d+", new_fv, r[i])
        self.__sync_cases(content, prod_id, version)
        return content

    def __to_obj(self, row: FtrRecordDoc, product: Product = None):
        obj = FtrRecordDocObj(**row.dict())
        obj.content = self.__autofill(self.__normalize_content(obj.content), row.product_id, product, row.version)
        if not (obj.file_no or "").strip():
            obj.file_no = self.__dhf_file_no(row.product_id)
        if product:
            obj.product_name = product.name
            obj.product_version = product.full_version
            obj.product_full_version = product.full_version
            obj.product_type_code = product.type_code
        obj.file_no = sync_file_no_version(obj.file_no or "", getattr(row, "version", None) or "") or obj.file_no
        return obj

    async def add_ftr_record_doc(self, form: FtrRecordDocForm):
        try:
            sql = select(func.count(FtrRecordDoc.id)).where(FtrRecordDoc.product_id == form.product_id, FtrRecordDoc.version == form.version)
            if db.session.execute(sql).scalar() > 0:
                return Resp.resp_err(msg=ts("msg_obj_exist"))
            row = FtrRecordDoc(**form.dict(exclude_none=True))
            row.id = None
            row.content = self.__normalize_content(row.content)
            db.session.add(row)
            db.session.commit()
            return Resp.resp_ok(data=FtrRecordDocForm(id=row.id))
        except Exception:
            logger.exception("")
            db.session.rollback()
        return Resp.resp_err(msg=ts(msg_err_db))

    async def duplicate_ftr_record_doc(self, id: int, product_id: int = None):
        try:
            fromdoc: FtrRecordDoc = db.session.execute(select(FtrRecordDoc).where(FtrRecordDoc.id == id)).scalars().first()
            if not fromdoc:
                return Resp.resp_err(msg=ts("msg_obj_null"))
            target_pid = product_id or fromdoc.product_id
            all_versions = db.session.execute(select(FtrRecordDoc.version).where(FtrRecordDoc.product_id == target_pid)).scalars().all()
            existing_set = {v for v in all_versions if v}
            if target_pid == fromdoc.product_id:
                version = new_version(fromdoc.version)
            else:
                def _seq(v):
                    m = re.search(r"(\d+)(?!.*\d)", v or "")
                    return int(m.group(1)) if m else -1
                valid = [v for v in all_versions if v]
                version = new_version(max(valid, key=_seq)) if valid else fromdoc.version
            while version in existing_set:
                version = new_version(version)
            newdoc = FtrRecordDoc(
                product_id=target_pid, version=version,
                file_no=sync_file_no_version((fromdoc.file_no or "").strip() or self.__dhf_file_no(target_pid), version) or None,
                change_log=fromdoc.change_log,
                content=copy.deepcopy(self.__normalize_content(fromdoc.content)),
            )
            db.session.add(newdoc)
            serv_review_util.add_raised_copy_dhf(
                target_pid, fromdoc.product_id, getattr(fromdoc, "file_no", None), version, DOC_NAME,
            )
            db.session.commit()
            if target_pid != fromdoc.product_id:
                await self.rebind_product(newdoc.id, target_pid)
            return Resp.resp_ok(data=FtrRecordDocForm(id=newdoc.id))
        except Exception:
            logger.exception("")
            db.session.rollback()
        return Resp.resp_err(msg=ts(msg_err_db))

    async def update_ftr_record_doc(self, form: FtrRecordDocForm):
        try:
            row: FtrRecordDoc = db.session.execute(select(FtrRecordDoc).where(FtrRecordDoc.id == form.id)).scalars().first()
            if not row:
                return Resp.resp_err(msg=ts("msg_obj_null"))
            for key, value in form.dict(exclude_none=True).items():
                if key == "id":
                    continue
                if key == "content":
                    value = self.__normalize_content(value)
                setattr(row, key, value)
            if getattr(form, "version", None) and hasattr(row, "file_no"):
                row.file_no = sync_file_no_version(row.file_no or "", form.version or row.version) or row.file_no
            db.session.commit()
            return Resp.resp_ok()
        except Exception:
            logger.exception("")
            db.session.rollback()
        return Resp.resp_err(msg=ts(msg_err_db))

    async def rebind_product(self, id: int, product_id: int):
        try:
            row: FtrRecordDoc = db.session.execute(select(FtrRecordDoc).where(FtrRecordDoc.id == id)).scalars().first()
            if not row:
                return Resp.resp_err(msg=ts("msg_obj_null"))
            product: Product = db.session.execute(select(Product).where(Product.id == product_id)).scalars().first()
            if not product:
                return Resp.resp_err(msg=ts("msg_obj_null"))
            db.session.execute(delete(FtrRecordDoc).where(FtrRecordDoc.product_id == product_id, FtrRecordDoc.version == row.version, FtrRecordDoc.id != id))
            content = copy.deepcopy(DEFAULT_CONTENT) if isinstance(DEFAULT_CONTENT, dict) else {"sections": []}
            content = self.__autofill(content, product_id, product, row.version, force=True)
            row.product_id = product_id
            row.content = content
            db.session.commit()
            return Resp.resp_ok(data=self.__to_obj(row, product))
        except Exception:
            logger.exception("")
            db.session.rollback()
        return Resp.resp_err(msg=ts(msg_err_db))

    async def delete_ftr_record_doc(self, id: int):
        db.session.execute(delete(FtrRecordDoc).where(FtrRecordDoc.id == id))
        db.session.commit()
        return Resp.resp_ok()

    async def get_ftr_record_doc(self, id: int):
        sql = select(FtrRecordDoc, Product).join(Product, FtrRecordDoc.product_id == Product.id).where(FtrRecordDoc.id == id)
        row = db.session.execute(sql).first()
        if not row:
            return Resp.resp_err(msg=ts("msg_obj_null"))
        doc, product = row
        return Resp.resp_ok(data=self.__to_obj(doc, product))

    async def list_ftr_record_doc(self, op_user: UserObj = None, product_id: int = 0, version: str = None, page_index: int = 0, page_size: int = 10):
        wheres = []
        if product_id:
            wheres.append(FtrRecordDoc.product_id == product_id)
        if version:
            wheres.append(FtrRecordDoc.version.like(f"%{version}%"))
        scope = product_scope_filter(op_user)
        if scope is not None:
            wheres.append(scope)
        sql_total = select(func.count(FtrRecordDoc.id)).join(Product, FtrRecordDoc.product_id == Product.id).where(*wheres)
        total = db.session.execute(sql_total).scalar() or 0
        sql = (select(FtrRecordDoc, Product).join(Product, FtrRecordDoc.product_id == Product.id).where(*wheres)
               .order_by(FtrRecordDoc.id.desc()).offset(page_index * page_size).limit(page_size))
        rows: List[FtrRecordDocObj] = [self.__to_obj(doc, product) for doc, product in db.session.execute(sql).all()]
        return Resp.resp_ok(data=Page(total=total, rows=rows, page_index=page_index, page_size=page_size))

    async def export_ftr_record_doc(self, output, id: int):
        resp = await self.get_ftr_record_doc(id)
        obj: FtrRecordDocObj = resp.data
        if obj is None:
            Document().save(output)
            output.seek(0)
            return
        content = obj.content if isinstance(obj.content, dict) else self.__normalize_content(obj.content)
        sections = content.get("sections") or []

        document = Document()
        section = document.sections[0]
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

        header_para = section.header.add_paragraph()
        header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        docx_util.fonted_txt(header_para, obj.file_no or "")
        docx_util.add_page_number_footer(section, obj.file_no or "")

        def set_cell(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
            s = str(text or "")
            cell.text = ""
            for i, line in enumerate(s.split("\n")):
                para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
                para.alignment = align
                para.paragraph_format.line_spacing = 1.3
                docx_util.fonted_txt(para, line, font_size=10.5, bold=bold)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        for s in sections:
            # 标题
            title = str(s.get("title") or "").strip()
            if title:
                p = document.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                docx_util.fonted_txt(p, title, font_size=18.0, bold=True)
                document.add_paragraph()
            # 表格
            for grid in (s.get("tables") or []):
                grid = [r for r in (grid or []) if isinstance(r, list)]
                cols = max((len(r) for r in grid), default=0)
                if cols <= 0:
                    continue
                # 现场测试记录大表固定 5 列，按 Word 模板还原合并单元格结构
                is_ftr_big_table = (cols == 5 and len(grid) >= 7)
                table = document.add_table(rows=0, cols=cols)
                table.style = "Table Grid"
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                rows_total = len(grid)
                for r_idx, r in enumerate(grid):
                    cells = table.add_row().cells
                    # 先清空所有单元格默认段落，避免 merge() 拼接文字
                    for c in cells:
                        c.text = ""
                    if not is_ftr_big_table:
                        for c_idx in range(cols):
                            set_cell(cells[c_idx], r[c_idx] if c_idx < len(r) else "")
                        continue
                    # 按模板合并：先做合并，再填充合并后单元格的值
                    if r_idx <= 4:
                        if r_idx in (0, 2):
                            cells[3].merge(cells[4])
                            spans = [(0, r[0]), (1, r[1]), (2, r[2]), (3, r[3])]
                        elif r_idx in (1, 3, 4):
                            cells[1].merge(cells[4])
                            spans = [(0, r[0]), (1, r[1])]
                    elif r_idx == 5:
                        cells[0].merge(cells[4])
                        spans = [(0, r[0])]
                    elif r_idx == 6:
                        cells[2].merge(cells[3])
                        spans = [(0, r[0]), (1, r[1]), (2, r[2]), (4, r[4])]
                    else:
                        cells[2].merge(cells[3])
                        spans = [(0, r[0]), (1, r[1]), (2, r[2]), (4, r[4])]
                    for c_idx, val in spans:
                        set_cell(cells[c_idx], val)
                # 纵向合并功能模块列(列0)：相同模块名连续行合并
                if is_ftr_big_table and rows_total > 7:
                    start = 7
                    prev_module = grid[7][0] if len(grid[7]) > 0 else ""
                    group_start = 7
                    for r_idx in range(8, rows_total):
                        cur_module = grid[r_idx][0] if len(grid[r_idx]) > 0 else ""
                        if cur_module != prev_module:
                            if r_idx - 1 > group_start:
                                # 先清空被合并行的列0文字，避免 merge() 拼接
                                for rr in range(group_start + 1, r_idx):
                                    table.cell(rr, 0).text = ""
                                table.cell(group_start, 0).merge(table.cell(r_idx - 1, 0))
                            group_start = r_idx
                            prev_module = cur_module
                    if rows_total - 1 > group_start:
                        for rr in range(group_start + 1, rows_total):
                            table.cell(rr, 0).text = ""
                        table.cell(group_start, 0).merge(table.cell(rows_total - 1, 0))

        document.save(output)
        output.seek(0)