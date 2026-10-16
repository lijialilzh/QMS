#!/usr/bin/env python
# encoding: utf-8

# 单元/集成/系统测试记录：导入用例后生成；导出按模版+用例多 sheet。

import copy
import logging
import os
import re
from collections import OrderedDict

from openpyxl import load_workbook
from sqlalchemy import select

from ..model.dev_test_record_doc import DevTestRecordDoc
from ..model.prod_dhf import ProdDhf
from ..model.product import Product
from ..model.test_case import TestCase
from ..model.test_set import TestSet
from ..utils.sql_ctx import db

logger = logging.getLogger(__name__)

STAGE_META = {
    "单元测试": {
        "title": "单元测试记录",
        "template": "unit_test_record_template.xlsx",
        "dhf_kw": "单元测试记录",
    },
    "集成测试": {
        "title": "集成测试记录",
        "template": "integration_test_record_template.xlsx",
        "dhf_kw": "集成测试记录",
    },
    "系统测试": {
        "title": "系统测试记录",
        "template": "system_test_record_template.xlsx",
        "dhf_kw": "系统测试记录",
    },
}

_TEMPLATE_DIR = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(__file__))),
    "src-res",
    "dev_test_record",
)

# 模版中用例表头关键字
_HEADER_KEYS = ("Case编号", "用例编号")


class Server(object):

    def __template_path(self, stage: str) -> str:
        meta = STAGE_META.get(stage) or {}
        return os.path.join(_TEMPLATE_DIR, meta.get("template") or "")

    def __dhf_file_no(self, prod_id, stage: str) -> str:
        kw = (STAGE_META.get(stage) or {}).get("dhf_kw") or stage
        if not prod_id:
            return ""
        row = db.session.execute(
            select(ProdDhf).where(ProdDhf.prod_id == prod_id, ProdDhf.name == kw).order_by(ProdDhf.id.asc())
        ).scalars().first()
        if not row:
            row = db.session.execute(
                select(ProdDhf).where(ProdDhf.prod_id == prod_id, ProdDhf.name.like(f"%{kw}%")).order_by(ProdDhf.id.asc())
            ).scalars().first()
        return (row.code or "").strip() if row and row.code else ""

    def __cases(self, prod_id, stage: str):
        ts_row = db.session.execute(
            select(TestSet).where(TestSet.product_id == prod_id, TestSet.stage == stage)
        ).scalars().first()
        if not ts_row:
            return []
        return list(db.session.execute(
            select(TestCase).where(TestCase.set_id == ts_row.id).order_by(TestCase.id.asc())
        ).scalars().all())

    async def ensure_from_case_import(self, product_id: int, stage: str):
        """单元/集成/系统用例导入成功后：无记录则建 A0，有则仅保证存在。"""
        if not product_id or stage not in STAGE_META:
            return
        try:
            existing = db.session.execute(
                select(DevTestRecordDoc).where(
                    DevTestRecordDoc.product_id == product_id,
                    DevTestRecordDoc.stage == stage,
                ).order_by(DevTestRecordDoc.id.desc())
            ).scalars().first()
            if existing:
                logger.info("dev_test_record exists id=%s stage=%s product_id=%s", existing.id, stage, product_id)
                return
            row = DevTestRecordDoc(
                product_id=product_id,
                stage=stage,
                version="A0",
                file_no=self.__dhf_file_no(product_id, stage) or None,
                content={"header": {}, "results": {}},
            )
            db.session.add(row)
            db.session.commit()
            logger.info("dev_test_record created id=%s stage=%s product_id=%s", row.id, stage, product_id)
        except Exception:
            logger.exception("ensure_from_case_import failed product_id=%s stage=%s", product_id, stage)
            db.session.rollback()

    def __find_header_row(self, ws):
        for r in range(1, min(120, (ws.max_row or 0) + 1)):
            v = ws.cell(r, 1).value
            if v and any(k in str(v) for k in _HEADER_KEYS):
                return r
        return None

    def __clear_data_rows(self, ws, header_row: int):
        if not header_row or not ws.max_row or ws.max_row <= header_row:
            return
        ws.delete_rows(header_row + 1, ws.max_row - header_row)

    def __write_cases(self, ws, header_row: int, cases, results: dict):
        """按模版 12 列写入用例行。"""
        for i, c in enumerate(cases):
            ridx = header_row + 1 + i
            code = (c.code or "").strip()
            res = (results or {}).get(code) or {}
            actual = res.get("actual")
            if actual is None or actual == "":
                actual = (c.expect or "")  # 默认实际结果同预期，可后改
            result = res.get("result") or ""
            remark = res.get("remark") or (c.note or "")
            vals = [
                code,
                c.srs_code or "",
                c.test_type or "",
                "",  # 阶段由调用方填
                c.function or "",
                c.description or "",
                c.precondition or "",
                c.test_step or "",
                c.expect or "",
                actual,
                result,
                remark,
            ]
            for col, val in enumerate(vals, 1):
                ws.cell(ridx, col, value=val)

    def export_record_xlsx(self, output, product_id: int, stage: str):
        """导出测试记录：按 sheet_name 分 sheet，布局用体系模版。"""
        meta = STAGE_META.get(stage)
        if not meta:
            raise ValueError(f"unsupported stage: {stage}")
        path = self.__template_path(stage)
        if not os.path.isfile(path):
            raise FileNotFoundError(path)

        record = db.session.execute(
            select(DevTestRecordDoc).where(
                DevTestRecordDoc.product_id == product_id,
                DevTestRecordDoc.stage == stage,
            ).order_by(DevTestRecordDoc.id.desc())
        ).scalars().first()
        results = {}
        if record and isinstance(record.content, dict):
            results = record.content.get("results") or {}

        product = db.session.execute(select(Product).where(Product.id == product_id)).scalars().first()
        cases = self.__cases(product_id, stage)
        groups = OrderedDict()
        for c in cases:
            sn = (getattr(c, "sheet_name", None) or "").strip() or meta["title"]
            groups.setdefault(sn, []).append(c)
        if not groups:
            groups[meta["title"]] = []

        wb = load_workbook(path)
        template = wb[wb.sheetnames[0]]
        header_row = self.__find_header_row(template) or 78

        def safe_title(name, used):
            base = re.sub(r'[\\/*?:\[\]]+', "_", name or meta["title"]).strip()[:31] or "Sheet1"
            title = base
            i = 2
            while title in used:
                suffix = f"_{i}"
                title = base[: 31 - len(suffix)] + suffix
                i += 1
            used.add(title)
            return title

        used = set()
        names = list(groups.keys())

        def fill_sheet(ws, group_cases):
            self.__clear_data_rows(ws, header_row)
            # 阶段列写入
            for c in group_cases:
                # monkey: stage on write
                pass
            for i, c in enumerate(group_cases):
                ridx = header_row + 1 + i
                code = (c.code or "").strip()
                res = results.get(code) or {}
                actual = res.get("actual")
                if actual is None or actual == "":
                    actual = c.expect or ""
                vals = [
                    code,
                    c.srs_code or "",
                    c.test_type or "",
                    stage,
                    c.function or "",
                    c.description or "",
                    c.precondition or "",
                    c.test_step or "",
                    c.expect or "",
                    actual,
                    res.get("result") or "",
                    res.get("remark") if res.get("remark") is not None else (c.note or ""),
                ]
                for col, val in enumerate(vals, 1):
                    ws.cell(ridx, col, value=val)
            # 文件编号 / 标题
            file_no = (record.file_no if record else None) or self.__dhf_file_no(product_id, stage)
            if file_no:
                ws.cell(1, 1, value=file_no)
            ws.cell(2, 1, value=meta["title"])
            if product and product.full_version:
                # 测试版本行尽量写入（模版约在 R75）
                for r in range(70, min(90, (ws.max_row or 80) + 1)):
                    if str(ws.cell(r, 1).value or "").strip() == "测试版本":
                        ws.cell(r, 2, value=product.full_version)
                        break

        fill_sheet(template, groups[names[0]])
        template.title = safe_title(names[0], used)
        for name in names[1:]:
            ws = wb.copy_worksheet(template)
            # copy 带了上一组数据，需再清再写
            hr = self.__find_header_row(ws) or header_row
            self.__clear_data_rows(ws, hr)
            # 重新写本组
            for i, c in enumerate(groups[name]):
                ridx = hr + 1 + i
                code = (c.code or "").strip()
                res = results.get(code) or {}
                actual = res.get("actual")
                if actual is None or actual == "":
                    actual = c.expect or ""
                vals = [
                    code, c.srs_code or "", c.test_type or "", stage,
                    c.function or "", c.description or "", c.precondition or "",
                    c.test_step or "", c.expect or "", actual,
                    res.get("result") or "",
                    res.get("remark") if res.get("remark") is not None else (c.note or ""),
                ]
                for col, val in enumerate(vals, 1):
                    ws.cell(ridx, col, value=val)
            ws.title = safe_title(name, used)

        wb.save(output)
        output.seek(0)
